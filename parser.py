"""Parse DeepBI exports as data, never as instructions or executable formulas."""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import math
import re
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

MAX_ROWS = 5000
MAX_COLUMNS = 50
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024

HEADERS = {
    "父ASIN": "asin", "商品名称": "name", "库存": "inventory",
    "总销售额": "total_sales", "广告销售额": "ad_sales", "广告总花费": "ad_spend",
    "自然销售额": "organic_sales", "自然销售占比": "organic_share", "TACOS": "tacos",
    "AI ACOS(SP)": "source_ai_acos", "AI广告花费(SP)": "ai_spend",
    "非AI ACOS(SP)": "nonai_acos", "SP广告销售额占比": "sp_sales_share",
    "AI广告销售额": "ai_sales", "AI销售额占比": "ai_sales_share",
    "AI订单量": "orders", "关键词种子数": "keyword_seeds", "ASIN流量种子数": "asin_seeds",
    "ACOS": "source_acos", "AI日预算": "daily_budget", "AI广告活动数量": "campaigns",
    "预期ACOS": "target", "托管策略": "strategy", "托管状态": "status", "托管时间": "managed_since",
}
REQUIRED = {"asin", "inventory", "total_sales", "ad_sales", "ad_spend", "orders", "target", "status"}
TEXT_FIELDS = {"asin", "name", "status", "strategy", "managed_since"}
RATE_FIELDS = {"organic_share", "tacos", "source_ai_acos", "nonai_acos", "sp_sales_share", "ai_sales_share", "source_acos", "target"}
INTEGER_FIELDS = {"inventory", "orders", "campaigns", "keyword_seeds", "asin_seeds"}
NEEDED_NUMBERS = {"inventory", "ad_sales", "ad_spend", "orders"}


class ReportParseError(ValueError):
    """A user-actionable export validation failure."""


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).replace("（", "(").replace("）", ")")


_KEYS = {_normalize_header(k): v for k, v in HEADERS.items()}


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value if value is None or isinstance(value, (str, int, float, bool)) else str(value)


def _numeric(value: Any, key: str, row: int, number_format: str = "") -> float | int | None:
    if value is None or isinstance(value, str) and value.strip() in {"", "-", "—", "--", "N/A", "n/a"}:
        if key in NEEDED_NUMBERS:
            raise ReportParseError(f"第{row}行的{key}缺少数值；请重新导出完整报表。")
        return None
    if isinstance(value, bool):
        raise ReportParseError(f"第{row}行的{key}不是有效数值。")
    if isinstance(value, (int, float)):
        result = float(value)
        if key in RATE_FIELDS and "%" in number_format:
            result *= 100
    else:
        # Only the leading value is read; comparison arrows are not current-period values.
        cleaned = str(value).strip().replace(",", "").replace("，", "")
        if key == "target":
            cleaned = re.sub(r"^自定义\s*", "", cleaned)
        match = re.match(r"^(?:[≤≥<>]=?\s*|(?:AED|SAR|USD)\s*|[$¥￥]\s*)?(-?\d+(?:\.\d+)?)\s*(?:%|％)?(?:\s|\(|（|$)", cleaned)
        if not match:
            raise ReportParseError(f"第{row}行的{key}无法解析：请保留DeepBI原始数字格式。")
        result = float(match.group(1))
    if not math.isfinite(result) or abs(result) > 1e14:
        raise ReportParseError(f"第{row}行的{key}数值超出可分析范围。")
    if key in NEEDED_NUMBERS | {"daily_budget", "campaigns", "keyword_seeds", "asin_seeds"} and result < 0:
        raise ReportParseError(f"第{row}行的{key}不能为负数。")
    if key in INTEGER_FIELDS:
        if not result.is_integer():
            raise ReportParseError(f"第{row}行的{key}应为整数。")
        return int(result)
    if key == "target" and not 16 <= result <= 50:
        raise ReportParseError(f"第{row}行目标ACOS应在16%至50%之间；例如24%或24。")
    return result


def parse_report(content: bytes) -> dict:
    """Return normalized rows and original cell values from one DeepBI worksheet."""
    if not content:
        raise ReportParseError("上传文件为空。")
    if len(content) > MAX_FILE_BYTES:
        raise ReportParseError("文件超过20MB，请导出单站点报表。")
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 2000 or sum(x.file_size for x in members) > MAX_UNCOMPRESSED_BYTES:
                raise ReportParseError("Excel解压后内容过大，请使用单站点原始导出。")
            if any(x.filename.lower().endswith("vbaproject.bin") for x in members):
                raise ReportParseError("不接受含宏的工作簿，请上传.xlsx原始导出。")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except ReportParseError:
        raise
    except Exception as exc:
        raise ReportParseError("无法读取Excel文件，请上传有效的.xlsx报表。") from exc

    try:
        candidates = []
        missing = REQUIRED
        for sheet in workbook.worksheets:
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                continue
            first = next(sheet.iter_rows(min_row=1, max_row=1, max_col=MAX_COLUMNS), ())
            cells = list(first)
            while cells and cells[-1].value is None:
                cells.pop()
            headers = [str(c.value or "").strip() for c in cells]
            keys = [_KEYS.get(_normalize_header(h)) for h in headers]
            absent = REQUIRED - set(keys)
            if len(absent) < len(missing):
                missing = absent
            if not absent:
                recognized = [k for k in keys if k]
                if len(set(recognized)) != len(recognized):
                    raise ReportParseError("报表含重复字段，请使用未经合并的原始导出。")
                candidates.append((sheet, headers, keys))
        if not candidates:
            friendly = [next((h for h, k in HEADERS.items() if k == field), field) for field in sorted(missing)]
            raise ReportParseError("未找到DeepBI报表，缺少字段：" + "、".join(friendly) + "。最多支持50列。")
        if len(candidates) > 1:
            raise ReportParseError("工作簿包含多张报表，请每次上传一个站点的一张报表。")
        sheet, headers, keys = candidates[0]
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
            raise ReportParseError("每份报表最多支持5000个Listing。")
        rows, seen, warnings = [], set(), []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(headers)), 2):
            if all(c.value is None or str(c.value).strip() == "" for c in cells):
                continue
            if len(rows) >= MAX_ROWS:
                raise ReportParseError("每份报表最多支持5000个Listing。")
            row: dict[str, Any] = {key: None for key in HEADERS.values()}
            row["source_row"] = row_number
            row["raw"] = {h: _json_value(c.value) for h, c in zip(headers, cells)}
            for key, cell in zip(keys, cells):
                if key is None:
                    continue
                if key in TEXT_FIELDS:
                    row[key] = str(cell.value).strip() if cell.value is not None else None
                else:
                    row[key] = _numeric(cell.value, key, row_number, cell.number_format or "")
            asin = (row["asin"] or "").upper()
            if not re.fullmatch(r"[A-Z0-9]{10}", asin):
                raise ReportParseError(f"第{row_number}行父ASIN格式不正确，应为10位字母或数字。")
            if asin in seen:
                raise ReportParseError(f"父ASIN {asin}重复出现；请先核对报表，避免重复计算。")
            seen.add(asin)
            row["asin"] = asin
            row["name"] = row["name"] or asin
            if not row["status"]:
                raise ReportParseError(f"第{row_number}行缺少托管状态。")
            if row["target"] is None:
                warnings.append(f"{asin}未导出单品目标，分析会标为目标缺失，不自动生成调参动作。")
            row["acos"] = round(row["ad_spend"] / row["ad_sales"] * 100, 6) if row["ad_sales"] > 0 else None
            rows.append(row)
        if not rows:
            raise ReportParseError("报表中没有可分析的Listing。")
        if any(key is None and header for key, header in zip(keys, headers)):
            warnings.append("存在未识别的附加列，原值已保存，但未用于本次规则计算。")
        return {"rows": rows, "warnings": warnings, "source_headers": headers, "sheet_name": sheet.title}
    except ReportParseError:
        raise
    except Exception as exc:
        raise ReportParseError("工作簿结构损坏或字段无法读取，请重新导出.xlsx报表。") from exc
    finally:
        workbook.close()
