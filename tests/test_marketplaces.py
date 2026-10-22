"""Market isolation, legacy migration and read-only language presentation checks."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
import importlib
import io
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from test_workflow import product, workbook_bytes


class MarketplaceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.import_store = tempfile.TemporaryDirectory(prefix='zhouce-market-module-')
        with patch.dict(os.environ, {
            'DEEPSEEK_API_KEY': '', 'WEEKLY_ACCESS_PASSWORD': '',
            'WEEKLY_REQUIRE_AUTH': '', 'WEEKLY_DATA_DIR': cls.import_store.name,
        }):
            cls.main = importlib.import_module('main')

    @classmethod
    def tearDownClass(cls):
        cls.import_store.cleanup()

    def setUp(self):
        self.env = patch.dict(os.environ, {
            'DEEPSEEK_API_KEY': '', 'WEEKLY_ACCESS_PASSWORD': '', 'WEEKLY_REQUIRE_AUTH': '',
        })
        self.env.start()
        self.temp = tempfile.TemporaryDirectory(prefix='zhouce-market-test-')
        self.store = Path(self.temp.name)
        self.client = TestClient(self.main.create_app(self.store))
        self.end = datetime.now(timezone.utc).date() - timedelta(days=1)
        self.start = self.end - timedelta(days=6)

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()
        self.env.stop()

    def upload(self, marketplace='AE', *, lang='zh', content=None, start=None, end=None):
        return self.client.post('/api/reports', params={'lang': lang}, data={
            'marketplace': marketplace,
            'start_date': (start or self.start).isoformat(),
            'end_date': (end or self.end).isoformat(),
        }, files={'file': ('weekly.xlsx', content or workbook_bytes(),
                           'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})

    def saved(self, marketplace='AE', **kwargs):
        response = self.upload(marketplace, **kwargs)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def confirm(self, saved, value, note):
        response = self.client.post(f"/api/reports/{saved['report']['id']}/confirmations", json={
            'analysis_id': saved['analysis']['id'],
            'executed_at': datetime.now(timezone.utc).isoformat(),
            'items': [{'asin': 'B000000001', 'actual_target': value, 'note': note}],
        })
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()['confirmations'][0]

    def settings(self, marketplace):
        response = self.client.get('/api/settings', params={'marketplace': marketplace})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def set_target(self, marketplace, value):
        response = self.client.put('/api/settings', params={'marketplace': marketplace},
                                   json={'target_acos': value, 'growth_priority': True})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def database_snapshots(self):
        with sqlite3.connect(self.store / 'weekly.sqlite3') as con:
            return con.execute('SELECT id,settings_json,result_json FROM analyses ORDER BY id').fetchall()

    def test_only_uae_and_saudi_markets_are_offered_with_separate_currencies(self):
        response = self.client.get('/api/marketplaces')
        self.assertEqual(response.status_code, 200, response.text)
        markets = {m['code']: m for m in response.json()['marketplaces']}
        self.assertEqual(set(markets), {'AE', 'SA'})
        self.assertEqual(markets['AE']['currency'], 'AED')
        self.assertEqual(markets['SA']['currency'], 'SAR')
        self.assertIn('阿联酋', markets['AE']['name_zh'])
        self.assertIn('沙特', markets['SA']['name_zh'])
        self.assertIn('Emirates', markets['AE']['name_en'])
        self.assertIn('Saudi', markets['SA']['name_en'])

    def test_same_asin_and_workbook_remain_separate_reports_in_each_market(self):
        content = workbook_bytes()
        ae, sa = self.saved('AE', content=content), self.saved('SA', content=content)
        self.assertNotEqual(ae['report']['id'], sa['report']['id'])
        self.assertFalse(ae['duplicate'])
        self.assertFalse(sa['duplicate'])
        for market, saved, currency in [('AE', ae, 'AED'), ('SA', sa, 'SAR')]:
            with self.subTest(market=market):
                reports = self.client.get('/api/reports', params={'marketplace': market}).json()['reports']
                self.assertEqual([r['id'] for r in reports], [saved['report']['id']])
                self.assertEqual(reports[0]['currency'], currency)
                self.assertEqual(saved['analysis']['report']['currency'], currency)
                duplicate = self.saved(market, content=content)
                self.assertTrue(duplicate['duplicate'])
                self.assertEqual(duplicate['report']['id'], saved['report']['id'])
        self.assertEqual(len(self.client.get('/api/reports').json()['reports']), 2)

    def test_confirmation_lists_and_history_are_market_specific_for_same_asin(self):
        ae, sa = self.saved('AE'), self.saved('SA')
        ae_event = self.confirm(ae, 20, '阿联酋实际调整')
        sa_event = self.confirm(sa, 28, 'Saudi decision')
        for market, event in [('AE', ae_event), ('SA', sa_event)]:
            events = self.client.get('/api/confirmations', params={'marketplace': market}).json()['confirmations']
            self.assertEqual([e['id'] for e in events], [event['id']])
            self.assertEqual(events[0]['note'], event['note'])
        self.assertEqual(len(self.client.get('/api/confirmations').json()['confirmations']), 2)
        for saved, target in [(ae, 20), (sa, 28)]:
            analyzed = self.client.post(f"/api/reports/{saved['report']['id']}/analyze").json()['analysis']
            history = analyzed['recommendations'][0]['history']['execution']
            self.assertEqual(history['actual_target'], target)
            self.assertEqual(history['report_id'], saved['report']['id'])

    def test_market_targets_apply_to_new_analysis_without_changing_other_market_or_history(self):
        self.set_target('AE', 20)
        self.assertEqual(self.settings('SA')['target_acos'], 24)
        self.set_target('SA', 32)
        ae, sa = self.saved('AE'), self.saved('SA')
        self.assertEqual(ae['analysis']['summary']['goal_acos'], 20)
        self.assertEqual(sa['analysis']['summary']['goal_acos'], 32)
        self.assertEqual(ae['analysis']['settings_snapshot']['target_acos'], 20)
        self.assertEqual(sa['analysis']['settings_snapshot']['target_acos'], 32)
        self.set_target('AE', 16)
        refreshed = self.client.post(f"/api/reports/{ae['report']['id']}/analyze").json()['analysis']
        self.assertEqual(refreshed['summary']['goal_acos'], 16)
        self.assertEqual(self.settings('SA')['target_acos'], 32)
        original = self.client.get(f"/api/reports/{ae['report']['id']}/analyses/{ae['analysis']['id']}").json()
        self.assertEqual(original['summary']['goal_acos'], 20)
        self.assertEqual(original['settings_snapshot']['target_acos'], 20)
        sa_again = self.client.get(f"/api/reports/{sa['report']['id']}/analysis").json()
        self.assertEqual(sa_again['id'], sa['analysis']['id'])
        self.assertEqual(sa_again['summary']['goal_acos'], 32)
        self.assertEqual(self.client.get('/api/settings').json()['target_acos'], 16)

    def test_legacy_global_settings_migrate_to_ae_without_leaking_to_new_sa_market(self):
        with tempfile.TemporaryDirectory(prefix='zhouce-legacy-settings-') as legacy:
            database = Path(legacy) / 'weekly.sqlite3'
            value = {'target_acos': 28.0, 'growth_priority': True, 'inventory_threshold': 3}
            with sqlite3.connect(database) as con:
                con.execute('CREATE TABLE settings (id INTEGER PRIMARY KEY CHECK(id=1), value TEXT NOT NULL)')
                con.execute('INSERT INTO settings VALUES(1,?)', (json.dumps(value),))
            with TestClient(self.main.create_app(legacy)) as upgraded:
                self.assertEqual(upgraded.get('/api/settings?marketplace=AE').json()['target_acos'], 28)
                self.assertEqual(upgraded.get('/api/settings?marketplace=SA').json()['target_acos'], 24)
                changed = upgraded.put('/api/settings?marketplace=AE',
                                       json={'target_acos': 20, 'growth_priority': True})
                self.assertEqual(changed.status_code, 200)
            with TestClient(self.main.create_app(legacy)) as reopened:
                self.assertEqual(reopened.get('/api/settings?marketplace=AE').json()['target_acos'], 20)
                self.assertEqual(reopened.get('/api/settings?marketplace=SA').json()['target_acos'], 24)

    def test_both_markets_records_and_settings_survive_application_restart(self):
        self.set_target('AE', 20)
        self.set_target('SA', 28)
        ae, sa = self.saved('AE'), self.saved('SA')
        self.confirm(sa, 24, '已执行')
        before = self.database_snapshots()
        self.client.close()
        self.client = TestClient(self.main.create_app(self.store))
        self.assertEqual(self.settings('AE')['target_acos'], 20)
        self.assertEqual(self.settings('SA')['target_acos'], 28)
        self.assertEqual(len(self.client.get('/api/reports').json()['reports']), 2)
        self.assertEqual(len(self.client.get('/api/confirmations?marketplace=SA').json()['confirmations']), 1)
        self.assertEqual(self.client.get('/api/confirmations?marketplace=AE').json()['confirmations'], [])
        self.assertEqual(self.database_snapshots(), before)
        for saved in (ae, sa):
            self.assertEqual(self.client.get(f"/api/reports/{saved['report']['id']}/analysis").json()['id'],
                             saved['analysis']['id'])

    def test_unsupported_markets_are_rejected_before_creating_reports_or_files(self):
        for marketplace in ['US', 'UK', 'DE', 'ZZ']:
            with self.subTest(marketplace=marketplace):
                self.assertEqual(self.upload(marketplace).status_code, 422)
                for endpoint in ['/api/settings', '/api/reports', '/api/confirmations']:
                    response = self.client.get(endpoint, params={'marketplace': marketplace})
                    self.assertEqual(response.status_code, 422, response.text)
                self.assertEqual(self.client.put('/api/settings', params={'marketplace': marketplace},
                    json={'target_acos': 20, 'growth_priority': True}).status_code, 422)
        self.assertEqual(self.client.get('/api/reports').json()['reports'], [])
        self.assertEqual(list((self.store / 'uploads').glob('*')), [])

    def test_unsupported_market_errors_follow_requested_english_language(self):
        responses = [self.upload('US', lang='en')]
        responses.extend(self.client.get(endpoint, params={'marketplace': 'US', 'lang': 'en'})
                         for endpoint in ['/api/settings', '/api/reports', '/api/confirmations'])
        responses.append(self.client.put('/api/settings', params={'marketplace': 'US', 'lang': 'en'},
                                         json={'target_acos': 20, 'growth_priority': True}))
        for response in responses:
            with self.subTest(url=str(response.request.url), method=response.request.method):
                self.assertEqual(response.status_code, 422, response.text)
                detail = response.json()['detail']
                self.assertIsInstance(detail, str)
                self.assertNotRegex(detail, r'[\u4e00-\u9fff]')
                self.assertIn('AE', detail)
                self.assertIn('SA', detail)
        self.assertEqual(self.client.get('/api/reports').json()['reports'], [])
        self.assertEqual(list((self.store / 'uploads').glob('*')), [])

    def test_saudi_currency_prefixed_amounts_parse_without_changing_financial_values(self):
        saved = self.saved('SA', content=workbook_bytes([product(**{
            '总销售额': 'SAR 3,000.00', '广告销售额': 'SAR 2,000.00',
            '广告总花费': 'SAR 500.00 (↑10%)', '自然销售额': 'SAR 1,000.00',
            'AI广告销售额': 'SAR 2,000.00', 'AI广告花费(SP)': 'SAR 500.00',
            'AI日预算': 'SAR 10.00',
        })]))
        self.assertEqual(saved['report']['currency'], 'SAR')
        summary = saved['analysis']['summary']
        self.assertEqual(summary['total_sales'], 3000)
        self.assertEqual(summary['ad_sales'], 2000)
        self.assertEqual(summary['ad_spend'], 500)
        self.assertEqual(summary['acos'], 25)
        self.assertTrue(summary['total_sales_reliable'])
        metrics = saved['analysis']['recommendations'][0]['metrics']
        self.assertEqual(metrics['spend'], 500)
        self.assertEqual(metrics['daily_budget'], 10)

    def test_language_switch_is_read_only_and_preserves_numbers_and_user_text(self):
        saved = self.saved('SA', content=workbook_bytes([product(商品名称='用户原始中文商品名', 库存=2)]))
        event = self.confirm(saved, 24, '用户原始备注不应被改写')
        before = self.database_snapshots()
        report_id, analysis_id = saved['report']['id'], saved['analysis']['id']
        zh = self.client.get(f'/api/reports/{report_id}/analysis?lang=zh').json()
        en = self.client.get(f'/api/reports/{report_id}/analysis?lang=en').json()
        english_version = self.client.get(f'/api/reports/{report_id}/analyses/{analysis_id}?lang=en').json()
        self.assertEqual(zh['id'], en['id'])
        self.assertEqual(en['id'], english_version['id'])
        self.assertEqual(en['created_at'], zh['created_at'])
        self.assertEqual(en['settings_snapshot'], zh['settings_snapshot'])
        for field in ['ad_spend', 'ad_sales', 'acos', 'orders', 'total_sales', 'goal_acos', 'listing_count']:
            self.assertEqual(en['summary'][field], zh['summary'][field])
        left, right = zh['recommendations'][0], en['recommendations'][0]
        for field in ['asin', 'name', 'current_target', 'recommended_target', 'action', 'actionable', 'metrics']:
            self.assertEqual(left[field], right[field])
        self.assertEqual(right['name'], '用户原始中文商品名')
        self.assertEqual(en['confirmations'][0]['note'], event['note'])
        self.assertNotEqual(zh['summary']['headline'], en['summary']['headline'])
        self.assertNotRegex(en['summary']['headline'], r'[\u4e00-\u9fff]')
        self.assertNotRegex(right['reason'], r'[\u4e00-\u9fff]')
        self.assertEqual(self.database_snapshots(), before)
        versions = self.client.get(f'/api/reports/{report_id}/analyses').json()['analyses']
        self.assertEqual(len(versions), 1)

    def test_english_upload_and_reanalysis_return_localized_views_but_keep_chinese_history(self):
        saved = self.saved('AE', lang='en')
        self.assertNotRegex(saved['analysis']['summary']['headline'], r'[\u4e00-\u9fff]')
        report_id = saved['report']['id']
        zh = self.client.get(f'/api/reports/{report_id}/analysis?lang=zh').json()
        self.assertRegex(zh['summary']['headline'], r'[\u4e00-\u9fff]')
        en = self.client.post(f'/api/reports/{report_id}/analyze?lang=en').json()['analysis']
        self.assertNotRegex(en['summary']['headline'], r'[\u4e00-\u9fff]')
        self.assertNotEqual(en['id'], saved['analysis']['id'])
        old = self.client.get(f"/api/reports/{report_id}/analyses/{saved['analysis']['id']}?lang=en").json()
        self.assertEqual(old['summary']['acos'], saved['analysis']['summary']['acos'])

    def test_english_export_uses_requested_snapshot_without_altering_saved_analysis(self):
        ai_summary = '这是原始中文AI复核，不应被伪装成翻译结果。'
        review = {'status': 'completed', 'provider': 'DeepSeek', 'model': 'deepseek-chat',
                  'summary': ai_summary, 'findings': [], 'limitations': [], 'error': None}
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': 'test-key-never-sent'}), \
                patch.object(self.main, 'review_analysis', new=AsyncMock(return_value=review)) as mocked:
            saved = self.saved('SA')
        mocked.assert_awaited_once()
        self.set_target('SA', 32)
        latest = self.client.post(f"/api/reports/{saved['report']['id']}/analyze").json()['analysis']
        self.assertNotEqual(latest['id'], saved['analysis']['id'])
        self.assertEqual(latest['summary']['goal_acos'], 32)
        before = self.database_snapshots()
        response = self.client.get(f"/api/reports/{saved['report']['id']}/export",
                                   params={'analysis_id': saved['analysis']['id'], 'lang': 'en'})
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        self.assertEqual(workbook.sheetnames, ['Overview', 'Recommendations'])
        self.assertEqual([c.value for c in workbook['Overview'][1]], ['Field', 'Value'])
        overview = dict(workbook['Overview'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(overview['Analysis version'], saved['analysis']['id'])
        self.assertEqual(overview['goal_acos'], 24)
        self.assertEqual(overview['Marketplace'], 'SA')
        self.assertEqual(overview['Currency'], 'SAR')
        self.assertEqual(overview['AI review summary (original language)'], ai_summary)
        self.assertEqual(overview['AI review status'], 'completed')
        headers = [c.value for c in workbook['Recommendations'][1]]
        self.assertTrue({'ASIN', 'Product', 'Current target', 'Suggested target', 'Reason', 'Ad ACOS'} <= set(headers))
        for header in headers:
            self.assertNotRegex(header, r'[\u4e00-\u9fff]')
        row = dict(zip(headers, (c.value for c in workbook['Recommendations'][2])))
        self.assertEqual(row['ASIN'], 'B000000001')
        self.assertEqual(row['Product'], '测试灯具')
        self.assertEqual(row['Current target'], 28)
        self.assertEqual(row['Suggested target'], saved['analysis']['recommendations'][0]['recommended_target'])
        self.assertEqual(row['Ad ACOS'], 25)
        self.assertEqual(row['Actionable'], 'No')
        self.assertNotRegex(row['Reason'], r'[\u4e00-\u9fff]')
        self.assertEqual(self.database_snapshots(), before)


if __name__ == '__main__':
    unittest.main()
