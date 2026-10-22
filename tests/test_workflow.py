"""Integration checks for the report → analysis → human confirmation workflow.

Run: python3 -m unittest discover -s tests -v
All tests use temporary SQLite stores and disable outbound DeepSeek calls.
"""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
import importlib
import io
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import AsyncMock, patch
import zipfile

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook


HEADERS = [
    '父ASIN', '商品名称', '库存', '总销售额', '广告销售额', '广告总花费',
    '自然销售额', '自然销售占比', 'TACOS', 'AI ACOS(SP)', 'AI广告花费(SP)',
    '非AI ACOS(SP)', 'SP广告销售额占比', 'AI广告销售额', 'AI销售额占比',
    'AI订单量', '关键词种子数', 'ASIN流量种子数', 'ACOS', 'AI日预算',
    'AI广告活动数量', '预期ACOS', '托管策略', '托管状态', '托管时间',
]


def product(**changes):
    row = dict(zip(HEADERS, [
        'B000000001', '测试灯具', 50, '300.00', '200.00', '50.00',
        '100.00', '33.33%', '16.67%', '25.00%', '50.00', '-', '100.00%',
        '200.00', '100.00%', '5', '12', '300', '25.00%', '10.00', '3',
        '≤ 28%', '平衡策略', '投放中', '26.05.28托管',
    ]))
    row.update(changes)
    return row


def workbook_bytes(products=None, headers=None):
    headers = headers or HEADERS
    book = Workbook()
    sheet = book.active
    sheet.title = 'AI托管列表'
    sheet.append(headers)
    for row in products if products is not None else [product()]:
        sheet.append([row.get(h) for h in headers])
    result = io.BytesIO()
    book.save(result)
    return result.getvalue()


class WeeklyWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.import_store = tempfile.TemporaryDirectory(prefix='weekly-module-test-')
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': '', 'WEEKLY_DATA_DIR': cls.import_store.name}):
            cls.main = importlib.import_module('main')

    @classmethod
    def tearDownClass(cls):
        cls.import_store.cleanup()

    def setUp(self):
        self.env = patch.dict(os.environ, {'DEEPSEEK_API_KEY': ''})
        self.env.start()
        self.temp = tempfile.TemporaryDirectory(prefix='weekly-workflow-test-')
        self.store = Path(self.temp.name)
        self.client = TestClient(self.main.create_app(self.store))
        self.end = datetime.now(timezone.utc).date() - timedelta(days=1)
        self.start = self.end - timedelta(days=6)

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()
        self.env.stop()

    def upload(self, content=None, start=None, end=None, marketplace='AE', filename='weekly.xlsx'):
        return self.client.post('/api/reports', data={
            'marketplace': marketplace,
            'start_date': (start or self.start).isoformat(),
            'end_date': (end or self.end).isoformat(),
        }, files={'file': (filename, content if content is not None else workbook_bytes(),
                           'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})

    def saved(self, **kwargs):
        response = self.upload(**kwargs)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def confirmation(self, saved, *, asin='B000000001', actual_target=24, note='已在DeepBI实际调整', executed_at=None):
        return {
            'analysis_id': saved['analysis']['id'],
            'executed_at': executed_at or datetime.now(timezone.utc).isoformat(),
            'items': [{'asin': asin, 'actual_target': actual_target, 'note': note}],
        }

    def post_confirmation(self, saved, **kwargs):
        return self.client.post(f"/api/reports/{saved['report']['id']}/confirmations", json=self.confirmation(saved, **kwargs))

    def test_upload_persists_original_and_normalized_analysis(self):
        content = workbook_bytes([product(广告总花费='50.00 (↑10%)')])
        data = self.saved(content=content)
        self.assertFalse(data['duplicate'])
        report, analysis = data['report'], data['analysis']
        self.assertEqual(report['row_count'], 1)
        self.assertEqual(report['day_count'], 7)
        self.assertEqual(analysis['summary']['ad_spend'], 50)
        self.assertEqual(analysis['summary']['ad_sales'], 200)
        self.assertAlmostEqual(analysis['summary']['acos'], 25)
        self.assertEqual(len(analysis['recommendations']), 1)
        self.assertEqual(self.client.get(f"/api/reports/{report['id']}/download").content, content)
        self.assertTrue((self.store / 'weekly.sqlite3').exists())
        self.assertEqual((self.store / 'uploads' / f"{report['id']}.xlsx").read_bytes(), content)

    def test_duplicate_upload_does_not_duplicate_report_or_snapshot(self):
        content = workbook_bytes()
        first, again = self.saved(content=content), self.saved(content=content)
        self.assertTrue(again['duplicate'])
        self.assertEqual(first['report']['id'], again['report']['id'])
        self.assertEqual(first['analysis']['id'], again['analysis']['id'])
        self.assertEqual(len(self.client.get('/api/reports').json()['reports']), 1)
        self.assertEqual(len(self.client.get(f"/api/reports/{first['report']['id']}/analyses").json()['analyses']), 1)

    def test_report_and_analysis_survive_process_recreation(self):
        original = self.saved()
        self.client.close()
        self.client = TestClient(self.main.create_app(self.store))
        reports = self.client.get('/api/reports').json()['reports']
        self.assertEqual(reports[0]['id'], original['report']['id'])
        restored = self.client.get(f"/api/reports/{reports[0]['id']}/analysis").json()
        self.assertEqual(restored['id'], original['analysis']['id'])
        self.assertEqual(restored['recommendations'], original['analysis']['recommendations'])

    def test_settings_and_reanalysis_preserve_old_snapshot(self):
        first = self.saved()
        changed = self.client.put('/api/settings', json={'target_acos': 20, 'growth_priority': True})
        self.assertEqual(changed.status_code, 200)
        new = self.client.post(f"/api/reports/{first['report']['id']}/analyze").json()['analysis']
        self.assertNotEqual(new['id'], first['analysis']['id'])
        self.assertEqual(new['settings_snapshot']['target_acos'], 20)
        old = self.client.get(f"/api/reports/{first['report']['id']}/analyses/{first['analysis']['id']}").json()
        self.assertEqual(old['settings_snapshot']['target_acos'], 24)
        self.assertEqual(old['summary'], first['analysis']['summary'])
        self.assertEqual(old['recommendations'], first['analysis']['recommendations'])

    def test_thirty_day_baseline_is_not_weekly_growth_comparator(self):
        previous_end = self.start - timedelta(days=1)
        self.saved(start=previous_end - timedelta(days=29), end=previous_end)
        current = self.saved()
        self.assertFalse(current['analysis']['comparison']['available'])
        metrics = current['analysis']['comparison'].get('metrics', {})
        self.assertIsNone(metrics.get('total_sales_pct'))
        self.assertIsNone(metrics.get('ad_sales_pct'))

    def test_adjacent_weekly_periods_have_valid_comparison(self):
        prior_end = self.start - timedelta(days=1)
        prior = self.saved(start=prior_end - timedelta(days=6), end=prior_end)
        current = self.saved(content=workbook_bytes([product(**{
            '总销售额': '600', '广告销售额': '400', '自然销售额': '200', '广告总花费': '80',
            'AI广告销售额': '400', 'AI广告花费(SP)': '80',
        })]))
        comparison = current['analysis']['comparison']
        self.assertTrue(comparison['available'], comparison)
        self.assertEqual(comparison['previous_report_id'], prior['report']['id'])
        self.assertAlmostEqual(comparison['metrics']['ad_sales_pct'], 100)

    def test_overlapping_or_other_market_periods_do_not_compare(self):
        self.saved(start=self.start - timedelta(days=2), end=self.end - timedelta(days=2))
        prior_end = self.start - timedelta(days=1)
        self.saved(start=prior_end - timedelta(days=6), end=prior_end, marketplace='SA')
        current = self.saved()
        self.assertFalse(current['analysis']['comparison']['available'])

    def test_persistent_cost_overrun_can_produce_a_confirmable_24_percent_trial(self):
        content = workbook_bytes([product(**{
            '广告总花费': '60', 'AI广告花费(SP)': '60',
            'ACOS': '30%', 'AI ACOS(SP)': '30%', 'TACOS': '20%',
        })])
        prior_end = self.start - timedelta(days=1)
        self.saved(content=content, start=prior_end - timedelta(days=6), end=prior_end)
        current = self.saved(content=content)
        recommendation = current['analysis']['recommendations'][0]
        self.assertTrue(recommendation['actionable'], recommendation)
        self.assertEqual(recommendation['current_target'], 28)
        self.assertEqual(recommendation['recommended_target'], 24)
        response = self.post_confirmation(current, actual_target=24, note='')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['confirmations'][0]['actual_target'], 24)

    def test_confirmed_action_is_used_only_for_its_market_and_correct_time_period(self):
        prior_end = self.start - timedelta(days=1)
        prior = self.saved(start=prior_end - timedelta(days=6), end=prior_end)
        self.assertEqual(self.post_confirmation(prior).status_code, 200)
        current = self.saved(content=workbook_bytes([product(预期ACOS='≤ 24%')]))
        recommendation = current['analysis']['recommendations'][0]
        execution = recommendation['history']['execution']
        self.assertTrue(execution['has_confirmation'])
        self.assertEqual(execution['actual_target'], 24)
        self.assertEqual(execution['report_id'], prior['report']['id'])
        self.assertFalse(execution['complete_post_change_period'])
        self.assertTrue(execution['blocks_new_change'])
        self.assertFalse(recommendation['actionable'])
        other_market = self.saved(marketplace='SA')
        other_execution = other_market['analysis']['recommendations'][0]['history']['execution']
        self.assertFalse(other_execution['has_confirmation'])

    def test_confirmation_keeps_actual_target_and_user_attestation(self):
        saved = self.saved()
        response = self.post_confirmation(saved, actual_target=20, note='实际改为20%，与建议不同，先做保守试验')
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()['confirmations'][0]
        self.assertEqual(item['actual_target'], 20)
        self.assertEqual(item['previous_target'], 28)
        self.assertEqual(item['recommended_target'], saved['analysis']['recommendations'][0]['recommended_target'])
        self.assertEqual(item['note'], '实际改为20%，与建议不同，先做保守试验')
        self.assertIn('executed_at', item)
        self.assertIn('confirmed_at', item)
        self.assertEqual(self.client.get('/api/confirmations').json()['confirmations'][0]['id'], item['id'])
        analysis = self.client.get(f"/api/reports/{saved['report']['id']}/analysis").json()
        self.assertEqual(analysis['recommendations'], saved['analysis']['recommendations'])
        self.assertEqual(analysis['confirmations'][0]['actual_target'], 20)

    def test_duplicate_confirmation_is_rejected_even_after_reanalysis(self):
        saved = self.saved()
        self.assertEqual(self.post_confirmation(saved).status_code, 200)
        self.assertEqual(self.post_confirmation(saved).status_code, 409)
        newer = self.client.post(f"/api/reports/{saved['report']['id']}/analyze").json()
        payload = self.confirmation(saved)
        payload['analysis_id'] = newer['analysis']['id']
        response = self.client.post(f"/api/reports/{saved['report']['id']}/confirmations", json=payload)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(len(self.client.get('/api/confirmations').json()['confirmations']), 1)

    def test_confirmation_rejects_unknown_asin_and_wrong_analysis(self):
        first = self.saved()
        second = self.saved(content=workbook_bytes([product(父ASIN='B000000002')]))
        self.assertEqual(self.post_confirmation(first, asin='B000000002').status_code, 422)
        payload = self.confirmation(first)
        payload['analysis_id'] = second['analysis']['id']
        response = self.client.post(f"/api/reports/{first['report']['id']}/confirmations", json=payload)
        self.assertEqual(response.status_code, 404)
        self.assertEqual(self.client.get('/api/confirmations').json()['confirmations'], [])

    def test_override_note_and_actual_execution_time_are_required(self):
        saved = self.saved()
        self.assertFalse(saved['analysis']['recommendations'][0]['actionable'])
        self.assertEqual(self.post_confirmation(saved, note='').status_code, 422)
        future = (datetime.now(timezone.utc) + timedelta(days=1)).isoformat()
        self.assertEqual(self.post_confirmation(saved, executed_at=future).status_code, 422)
        self.assertEqual(self.post_confirmation(saved, executed_at='2020-01-01T00:00:00Z').status_code, 422)
        self.assertEqual(self.post_confirmation(saved, executed_at=datetime.now().isoformat()).status_code, 422)
        self.assertEqual(self.client.get('/api/confirmations').json()['confirmations'], [])

    def test_execution_before_selected_analysis_is_not_attributed_to_it(self):
        before_analysis = datetime.now(timezone.utc)
        saved = self.saved()
        response = self.post_confirmation(saved, executed_at=before_analysis.isoformat())
        self.assertEqual(response.status_code, 422)
        self.assertEqual(self.client.get('/api/confirmations').json()['confirmations'], [])

    def test_invalid_workbooks_leave_no_report_or_upload(self):
        bad = self.upload(content=b'not a workbook')
        self.assertEqual(bad.status_code, 422)
        missing = self.upload(content=workbook_bytes(headers=[h for h in HEADERS if h != '广告总花费']))
        self.assertEqual(missing.status_code, 422)
        duplicate = self.upload(content=workbook_bytes([product(), product()]))
        self.assertEqual(duplicate.status_code, 422)
        self.assertEqual(self.client.get('/api/reports').json()['reports'], [])
        self.assertEqual(list((self.store / 'uploads').glob('*')), [])

    def test_zero_ad_sales_never_becomes_zero_acos_success(self):
        saved = self.saved(content=workbook_bytes([product(广告销售额='0', AI广告销售额='0',
                    AI订单量='0', ACOS='0.00%', 总销售额='100', 自然销售额='100')]))
        analysis = saved['analysis']
        self.assertIsNone(analysis['summary']['acos'])
        recommendation = analysis['recommendations'][0]
        self.assertIsNone(recommendation['metrics']['acos'])
        self.assertFalse(recommendation['actionable'])
        self.assertNotEqual(recommendation['action'], 'expand')
        self.assertEqual(analysis['summary']['zero_sale_spend'], 50)

    def test_stock_protection_applies_strictly_below_three(self):
        saved = self.saved(content=workbook_bytes([
            product(父ASIN='B000000001', 库存=2),
            product(父ASIN='B000000002', 库存=3),
        ]))
        by_asin = {r['asin']: r for r in saved['analysis']['recommendations']}
        self.assertEqual(by_asin['B000000001']['action'], 'protect_inventory')
        self.assertFalse(by_asin['B000000001']['actionable'])
        self.assertNotEqual(by_asin['B000000002']['action'], 'protect_inventory')

    def test_unconfigured_deepseek_is_explicit_and_no_secret_is_returned(self):
        saved = self.saved()
        self.assertEqual(saved['analysis']['mode'], 'rules')
        self.assertEqual(saved['analysis']['ai_review']['status'], 'unavailable')
        settings = self.client.get('/api/settings').json()
        self.assertFalse(settings['ai']['configured'])
        # A helpful message may name DEEPSEEK_API_KEY; a credential field/value must not be exposed.
        self.assertFalse({'api_key', 'key', 'secret', 'token'} & set(settings['ai']))
        self.assertFalse(self.client.get('/api/health').json()['ai_configured'])

    def test_mock_ai_review_is_persisted_without_changing_rule_targets(self):
        review = {'status': 'completed', 'provider': 'DeepSeek', 'model': 'deepseek-chat',
                  'generated_at': datetime.now(timezone.utc).isoformat(), 'summary': '需要更多同周期数据。',
                  'findings': [{'asin': 'B000000001', 'comment': '保持观察。', 'possible_causes': [], 'next_steps': []}],
                  'limitations': [], 'error': None}
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': 'test-key-never-sent'}), \
                patch.object(self.main, 'review_analysis', new=AsyncMock(return_value=review)) as mocked:
            saved = self.saved()
            result = self.client.get(f"/api/reports/{saved['report']['id']}/analysis").json()
        mocked.assert_awaited_once()
        self.assertEqual(result['ai_review']['status'], 'completed')
        self.assertEqual(result['ai_review']['summary'], review['summary'])
        self.assertEqual(result['recommendations'], saved['analysis']['recommendations'])

    def test_mock_ai_failure_preserves_original_report_and_rule_analysis(self):
        content = workbook_bytes()
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': 'test-key-never-sent'}), \
                patch.object(self.main, 'review_analysis', new=AsyncMock(side_effect=RuntimeError('private upstream detail'))):
            saved = self.saved(content=content)
            result = self.client.get(f"/api/reports/{saved['report']['id']}/analysis").json()
        self.assertEqual(result['ai_review']['status'], 'failed')
        self.assertNotIn('private upstream detail', result['ai_review']['error'])
        self.assertEqual(result['recommendations'], saved['analysis']['recommendations'])
        self.assertEqual(self.client.get(f"/api/reports/{saved['report']['id']}/download").content, content)

    def test_backup_restores_reports_snapshots_and_confirmations(self):
        saved = self.saved()
        self.assertEqual(self.post_confirmation(saved, actual_target=20).status_code, 200)
        response = self.client.get('/api/backup')
        self.assertEqual(response.status_code, 200)
        with tempfile.TemporaryDirectory(prefix='weekly-restore-test-') as destination:
            with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                self.assertIn('weekly.sqlite3', archive.namelist())
                self.assertIn(f"uploads/{saved['report']['id']}.xlsx", archive.namelist())
                self.assertFalse(any('.env' in name for name in archive.namelist()))
                archive.extractall(destination)
            with TestClient(self.main.create_app(destination)) as restored:
                reports = restored.get('/api/reports').json()['reports']
                self.assertEqual(reports[0]['id'], saved['report']['id'])
                self.assertEqual(restored.get('/api/confirmations').json()['confirmations'][0]['actual_target'], 20)
                self.assertEqual(restored.get(f"/api/reports/{reports[0]['id']}/analysis").json()['id'], saved['analysis']['id'])

    def test_export_includes_analysis_and_does_not_emit_zero_sales_acos(self):
        saved = self.saved(content=workbook_bytes([product(广告销售额='0', AI广告销售额='0', AI订单量='0',
                    ACOS='0.00%', 总销售额='100', 自然销售额='100')]))
        response = self.client.get(f"/api/reports/{saved['report']['id']}/export")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        self.assertIn('分析概况', workbook.sheetnames)
        sheet = workbook['逐品建议']
        headers = [c.value for c in sheet[1]]
        self.assertEqual(sheet.cell(2, 1).value, 'B000000001')
        self.assertIsNone(sheet.cell(2, headers.index('广告ACOS') + 1).value)

    def test_export_preserves_selected_historical_analysis_and_rejects_foreign_version(self):
        first = self.saved()
        report_id = first['report']['id']
        old_analysis_id = first['analysis']['id']
        settings = self.client.put('/api/settings', json={'target_acos': 20, 'growth_priority': True})
        self.assertEqual(settings.status_code, 200)
        latest = self.client.post(f'/api/reports/{report_id}/analyze').json()['analysis']

        def overview(response):
            self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
            workbook = load_workbook(io.BytesIO(response.content), data_only=True)
            return dict(workbook['分析概况'].iter_rows(min_row=2, values_only=True))

        old_export = overview(self.client.get(f'/api/reports/{report_id}/export',
                                              params={'analysis_id': old_analysis_id}))
        latest_export = overview(self.client.get(f'/api/reports/{report_id}/export'))
        self.assertEqual(old_export['分析版本'], old_analysis_id)
        self.assertEqual(old_export['goal_acos'], 24)
        self.assertEqual(latest_export['分析版本'], latest['id'])
        self.assertEqual(latest_export['goal_acos'], 20)
        foreign = self.saved(content=workbook_bytes([product(父ASIN='B000000002')]))
        for invalid_id in ['missing-analysis-version', foreign['analysis']['id']]:
            with self.subTest(analysis_id=invalid_id):
                response = self.client.get(f'/api/reports/{report_id}/export', params={'analysis_id': invalid_id})
                self.assertEqual(response.status_code, 404)

    def test_cross_origin_mutation_is_rejected(self):
        response = self.client.put('/api/settings', json={'target_acos': 20, 'growth_priority': True},
                                   headers={'Origin': 'https://unrelated.example'})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.get('/api/settings').json()['target_acos'], 24)


if __name__ == '__main__':
    unittest.main()
